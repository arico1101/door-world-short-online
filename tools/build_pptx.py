# -*- coding: utf-8 -*-
import json, io, os, re
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from copy import deepcopy
from lxml import etree

PX = 9525                      # 96dpi の 1px
W, H = 1440, 811
FONT = "Zen Maru Gothic"
PAGE_BG = "F2F9FD"
HL = "FCE69F"                  # 黄色いマーカー（0.65の透過を白地に重ねた近似）
SCR = os.environ.get('TOBIRA_WORK', os.path.dirname(os.path.abspath(__file__)))

def px(v): return Emu(int(round(v * PX)))
def pt(v): return Pt(v * 0.75)

def rgb(s):
    m = re.match(r'rgba?\(([\d.]+),\s*([\d.]+),\s*([\d.]+)(?:,\s*([\d.]+))?\)', s or '')
    if not m: return None, 1.0
    a = float(m.group(4)) if m.group(4) is not None else 1.0
    return RGBColor(int(float(m.group(1))), int(float(m.group(2))), int(float(m.group(3)))), a

def set_font(run, size_px, color, bold):
    f = run.font
    f.size = pt(size_px)
    f.bold = bold
    c, _ = rgb(color)
    if c is not None: f.color.rgb = c
    f.name = FONT
    # 日本語は東アジア用のフォント指定も要る
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:ea', 'a:cs'):
        el = rPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag.split(':')[1])
        if el is None:
            el = etree.SubElement(rPr, '{http://schemas.openxmlformats.org/drawingml/2006/main}' + tag.split(':')[1])
        el.set('typeface', FONT)

ALIGN = {'left': PP_ALIGN.LEFT, 'start': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
         'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY}

def load_notes(repo):
    """SLIDES_SCRIPT.md をページごとに切り出して、発表者ノート用の素の文にする。"""
    path = os.path.join(repo, 'SLIDES_SCRIPT.md')
    if not os.path.exists(path):
        return {}
    src = io.open(path, encoding='utf-8').read()
    parts = re.split(r'^## (\d+)\. (.+)$', src, flags=re.M)
    notes = {}
    for i in range(1, len(parts), 3):
        num, title, body = int(parts[i]), parts[i + 1].strip(), parts[i + 2]
        lines = []
        for line in body.splitlines():
            line = line.rstrip()
            if line.strip() in ('---', ''):
                lines.append('')
                continue
            if line.startswith('# '):          # 次のパートの見出しは入れない
                continue
            line = re.sub(r'^>\s?', '', line)   # 読み上げの引用記号を外す
            line = re.sub(r'\*\*(.+?)\*\*', r'\1', line)
            line = re.sub(r'`(.+?)`', r'\1', line)
            lines.append(line)
        text = re.sub(r'\n{3,}', '\n\n', '\n'.join(lines)).strip()
        notes[num] = ('%d. %s\n\n%s' % (num, title, text))[:4800]
    return notes


def build(data, out):
    prs = Presentation()
    prs.slide_width, prs.slide_height = px(W), px(H)
    blank = prs.slide_layouts[6]
    imap = json.load(io.open(os.path.join(SCR, 'icon_map.json'), encoding='utf-8'))
    notes = load_notes(os.environ.get('TOBIRA_REPO', os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

    for page in data['pages']:
        sl = prs.slides.add_slide(blank)
        note = notes.get(page['p'])
        if note:
            sl.notes_slide.notes_text_frame.text = note
        bg = sl.background.fill; bg.solid(); bg.fore_color.rgb = RGBColor.from_string(PAGE_BG)
        items = page['items']
        order = ([i for i in items if i['k'] == 'box'] + [i for i in items if i['k'] == 'hl'] +
                 [i for i in items if i['k'] == 'img'] + [i for i in items if i['k'] == 'text'] +
                 [i for i in items if i['k'] == 'icon'])
        for it in order:
            k = it['k']
            if k in ('box', 'hl'):
                if k == 'hl':
                    bar = 5.7
                    sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(it['l']),
                                             px(it['t'] + it['h'] - bar), px(it['w']), px(bar))
                    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(HL)
                    sh.line.fill.background(); sh.shadow.inherit = False
                    continue
                rad = it.get('rad', '0px')
                w, h = it['w'], it['h']
                if '%' in rad and abs(w - h) < 2:
                    sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, px(it['l']), px(it['t']), px(w), px(h))
                else:
                    r = float(re.sub(r'[^\d.]', '', rad) or 0)
                    if r >= min(w, h) / 2 - 0.5 and r > 4:
                        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(it['l']), px(it['t']), px(w), px(h))
                        sh.adjustments[0] = 0.5
                    elif r > 0.5:
                        sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px(it['l']), px(it['t']), px(w), px(h))
                        sh.adjustments[0] = min(0.5, r / min(w, h))
                    else:
                        sh = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, px(it['l']), px(it['t']), px(w), px(h))
                c, a = rgb(it.get('bg'))
                if c is None or a == 0: sh.fill.background()
                else: sh.fill.solid(); sh.fill.fore_color.rgb = c
                bc, ba = rgb(it.get('br'))
                if bc is not None and it.get('bwd', 0) > 0 and ba > 0:
                    sh.line.color.rgb = bc; sh.line.width = px(it['bwd'])
                else:
                    sh.line.fill.background()
                sh.shadow.inherit = False
                if sh.has_text_frame: sh.text_frame.text = ''
            elif k == 'img':
                p = os.path.join(os.environ.get('TOBIRA_REPO', os.path.dirname(os.path.dirname(os.path.abspath(__file__)))), it['src'])
                if os.path.exists(p):
                    l, t, w, h = it['l'], it['t'], it['w'], it['h']
                    nw, nh = it.get('nw') or 0, it.get('nh') or 0
                    # HTML は object-fit:contain なので、箱いっぱいに伸ばさず実寸比で収める
                    if it.get('fit') == 'contain' and nw > 0 and nh > 0:
                        sc = min(w / nw, h / nh)
                        fw, fh = nw * sc, nh * sc
                        top_align = 'top' in (it.get('pos') or '')
                        l += (w - fw) / 2
                        t += 0 if top_align else (h - fh) / 2
                        w, h = fw, fh
                    sl.shapes.add_picture(p, px(l), px(t), px(w), px(h))
            elif k == 'icon':
                key = '%s|%s' % (it['ref'].lstrip('#'), it['col'])
                f = imap.get(key)
                if f:
                    p = os.path.join(SCR, 'icons', f + '.png')
                    if os.path.exists(p):
                        sl.shapes.add_picture(p, px(it['l']), px(it['t']), px(it['w']), px(it['h']))
            elif k == 'text':
                al = it.get('align', 'left')
                l, w = it['l'], it['w']
                # Canva側の字幅がわずかに違うので、折り返しが境界で暴れないよう余白を足す
                if al in ('center',): l -= 30; w += 60
                elif al in ('right', 'end'): l -= 36; w += 36
                else: w += 36
                tb = sl.shapes.add_textbox(px(l), px(it['t']), px(w), px(it['h'] + 6))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.TOP
                para = tf.paragraphs[0]
                para.alignment = ALIGN.get(al, PP_ALIGN.LEFT)
                lh = it.get('lh')
                try:
                    if lh and lh.endswith('px'): para.line_spacing = pt(float(lh[:-2]))
                except Exception: pass
                def new_para():
                    p = tf.add_paragraph()
                    p.alignment = ALIGN.get(al, PP_ALIGN.LEFT)
                    try:
                        if lh and lh.endswith('px'): p.line_spacing = pt(float(lh[:-2]))
                    except Exception: pass
                    return p
                for r in it['runs']:
                    if r.get('br'):
                        para = new_para(); continue
                    run = para.add_run()
                    run.text = r['s']
                    try: fw = int(r.get('fw', 400))
                    except Exception: fw = 700
                    set_font(run, r.get('fs', it['fs']), r.get('c'), fw >= 600)
    prs.save(out)
    return out

data = json.load(io.open(os.path.join(SCR, 'full.json'), encoding='utf-8'))
out = build(data, os.path.join(SCR, 'slides.pptx'))
print('書き出し:', out, os.path.getsize(out), 'bytes')
